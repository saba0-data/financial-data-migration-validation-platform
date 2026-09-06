import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference


# ============================================================
# PROJECT PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

RAW_DATA_DIR = BASE_DIR / "data" / "raw"
CLEANED_DATA_DIR = BASE_DIR / "data" / "cleaned"
REPORTS_DIR = BASE_DIR / "reports"

REPORTS_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# REPORT FILE PATHS
# ============================================================

DATA_QUALITY_REPORT = (
    REPORTS_DIR / "data_quality_report.csv"
)

RCA_REPORT = (
    REPORTS_DIR / "root_cause_analysis_report.csv"
)

POST_LOAD_REPORT = (
    REPORTS_DIR / "post_load_validation_report.csv"
)

SQL_VALIDATION_REPORT = (
    REPORTS_DIR / "sql_validation_report.csv"
)

OUTPUT_FILE = (
    REPORTS_DIR / "Financial_Data_Quality_Dashboard.xlsx"
)


# ============================================================
# LOAD CSV SAFELY
# ============================================================

def load_csv(file_path):

    if file_path.exists():

        return pd.read_csv(file_path)

    else:

        print(f"Warning: File not found -> {file_path}")

        return pd.DataFrame()


# ============================================================
# LOAD REPORTS
# ============================================================

def load_reports():

    data_quality_df = load_csv(
        DATA_QUALITY_REPORT
    )

    rca_df = load_csv(
        RCA_REPORT
    )

    post_load_df = load_csv(
        POST_LOAD_REPORT
    )

    sql_validation_df = load_csv(
        SQL_VALIDATION_REPORT
    )

    return (
        data_quality_df,
        rca_df,
        post_load_df,
        sql_validation_df
    )


# ============================================================
# CALCULATE SOURCE RECORDS
# ============================================================

def calculate_source_records():

    total_records = 0

    source_counts = {}

    files = {

        "customers": "customers.csv",

        "accounts": "accounts.csv",

        "transactions": "transactions.csv"

    }


    for table_name, file_name in files.items():

        file_path = (
            RAW_DATA_DIR / file_name
        )

        dataframe = load_csv(
            file_path
        )

        record_count = len(
            dataframe
        )

        source_counts[
            table_name
        ] = record_count

        total_records += record_count


    return (
        total_records,
        source_counts
    )


# ============================================================
# CALCULATE VALID / CLEANED RECORDS
# ============================================================

def calculate_valid_records():

    total_valid = 0

    valid_counts = {}

    files = {

        "customers": "customers_cleaned.csv",

        "accounts": "accounts_cleaned.csv",

        "transactions": "transactions_cleaned.csv"

    }


    for table_name, file_name in files.items():

        file_path = (
            CLEANED_DATA_DIR / file_name
        )

        dataframe = load_csv(
            file_path
        )

        record_count = len(
            dataframe
        )

        valid_counts[
            table_name
        ] = record_count

        total_valid += record_count


    return (
        total_valid,
        valid_counts
    )


# ============================================================
# CALCULATE DASHBOARD METRICS
# ============================================================

def calculate_metrics(
    data_quality_df,
    rca_df,
    post_load_df,
    sql_validation_df
):

    # --------------------------------------------------------
    # SOURCE RECORDS
    # --------------------------------------------------------

    (
        total_source_records,
        source_counts
    ) = calculate_source_records()


    # --------------------------------------------------------
    # VALID RECORDS
    # --------------------------------------------------------

    (
        total_valid_records,
        valid_counts
    ) = calculate_valid_records()


    # --------------------------------------------------------
    # INVALID RECORDS
    # --------------------------------------------------------

    total_invalid_records = (
        total_source_records -
        total_valid_records
    )


    # --------------------------------------------------------
    # MIGRATION SUCCESS RATE
    # --------------------------------------------------------

    if total_source_records > 0:

        migration_success_rate = (
            total_valid_records /
            total_source_records
        ) * 100

    else:

        migration_success_rate = 0


    # --------------------------------------------------------
    # DATA QUALITY ISSUES
    # --------------------------------------------------------

    total_issues = len(
        data_quality_df
    )


    if total_source_records > 0:

        issue_rate = (
            total_issues /
            total_source_records
        ) * 100

    else:

        issue_rate = 0


    # --------------------------------------------------------
    # PRIORITY COUNTS
    # --------------------------------------------------------

    high_priority = 0

    medium_priority = 0

    low_priority = 0


    if (
        not rca_df.empty
        and "priority" in rca_df.columns
    ):

        high_priority = len(
            rca_df[
                rca_df["priority"]
                .astype(str)
                .str.lower() == "high"
            ]
        )


        medium_priority = len(
            rca_df[
                rca_df["priority"]
                .astype(str)
                .str.lower() == "medium"
            ]
        )


        low_priority = len(
            rca_df[
                rca_df["priority"]
                .astype(str)
                .str.lower() == "low"
            ]
        )


    # --------------------------------------------------------
    # POST-LOAD VALIDATION
    # --------------------------------------------------------

    post_load_total = len(
        post_load_df
    )

    post_load_passed = 0


    if (
        not post_load_df.empty
        and "status" in post_load_df.columns
    ):

        post_load_passed = len(
            post_load_df[
                post_load_df["status"]
                .astype(str)
                .str.upper() == "PASS"
            ]
        )


    if post_load_total > 0:

        post_load_success_rate = (
            post_load_passed /
            post_load_total
        ) * 100

    else:

        post_load_success_rate = 0


    # --------------------------------------------------------
    # SQL VALIDATION
    # --------------------------------------------------------

    sql_total = len(
        sql_validation_df
    )

    sql_passed = 0


    if (
        not sql_validation_df.empty
        and "status" in sql_validation_df.columns
    ):

        sql_passed = len(
            sql_validation_df[
                sql_validation_df["status"]
                .astype(str)
                .str.upper() == "PASS"
            ]
        )


    if sql_total > 0:

        sql_success_rate = (
            sql_passed /
            sql_total
        ) * 100

    else:

        sql_success_rate = 0


    # --------------------------------------------------------
    # RETURN ALL METRICS
    # --------------------------------------------------------

    return {

        "total_source_records":
            total_source_records,

        "total_valid_records":
            total_valid_records,

        "total_invalid_records":
            total_invalid_records,

        "migration_success_rate":
            migration_success_rate,

        "total_issues":
            total_issues,

        "issue_rate":
            issue_rate,

        "high_priority":
            high_priority,

        "medium_priority":
            medium_priority,

        "low_priority":
            low_priority,

        "post_load_passed":
            post_load_passed,

        "post_load_total":
            post_load_total,

        "post_load_success_rate":
            post_load_success_rate,

        "sql_passed":
            sql_passed,

        "sql_total":
            sql_total,

        "sql_success_rate":
            sql_success_rate,

        "source_counts":
            source_counts,

        "valid_counts":
            valid_counts

    }


# ============================================================
# ADD DATAFRAME TO EXCEL SHEET
# ============================================================

def add_dataframe_to_sheet(
    worksheet,
    dataframe
):

    if dataframe.empty:

        worksheet["A1"] = (
            "No data available"
        )

        return


    # --------------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )


    header_font = Font(
        bold=True,
        color="FFFFFF"
    )


    # --------------------------------------------------------
    # WRITE HEADERS
    # --------------------------------------------------------

    for column_index, column_name in enumerate(
        dataframe.columns,
        start=1
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name
        )

        cell.font = header_font

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center"
        )


    # --------------------------------------------------------
    # WRITE DATA
    # --------------------------------------------------------

    for row_index, row in enumerate(
        dataframe.itertuples(index=False),
        start=2
    ):

        for column_index, value in enumerate(
            row,
            start=1
        ):

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )


    # --------------------------------------------------------
    # AUTO ADJUST COLUMN WIDTH
    # --------------------------------------------------------

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = (
            column_cells[0]
            .column_letter
        )


        for cell in column_cells:

            try:

                value_length = len(
                    str(cell.value)
                )

                if value_length > max_length:

                    max_length = value_length

            except Exception:

                pass


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            50
        )


# ============================================================
# CREATE DASHBOARD
# ============================================================

def create_dashboard(
    workbook,
    metrics,
    data_quality_df,
    rca_df
):

    worksheet = workbook.active

    worksheet.title = "Dashboard"


    # ========================================================
    # TITLE
    # ========================================================

    worksheet.merge_cells(
        "A1:I1"
    )


    title_cell = worksheet["A1"]


    title_cell.value = (
        "FINANCIAL DATA MIGRATION & QUALITY DASHBOARD"
    )


    title_cell.font = Font(
        bold=True,
        size=16,
        color="FFFFFF"
    )


    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )


    title_cell.alignment = Alignment(
        horizontal="center"
    )


    # ========================================================
    # DASHBOARD METRICS
    # ========================================================

    metrics_data = [

        (
            "Total Source Records",
            metrics[
                "total_source_records"
            ]
        ),

        (
            "Valid Records",
            metrics[
                "total_valid_records"
            ]
        ),

        (
            "Invalid Records",
            metrics[
                "total_invalid_records"
            ]
        ),

        (
            "Migration Success Rate",
            f"{metrics['migration_success_rate']:.2f}%"
        ),

        (
            "Data Quality Issues",
            metrics[
                "total_issues"
            ]
        ),

        (
            "Data Quality Issue Rate",
            f"{metrics['issue_rate']:.2f}%"
        ),

        (
            "High Priority Issues",
            metrics[
                "high_priority"
            ]
        ),

        (
            "Medium Priority Issues",
            metrics[
                "medium_priority"
            ]
        ),

        (
            "Low Priority Issues",
            metrics[
                "low_priority"
            ]
        ),

        (
            "Post-Load Validation",
            f"{metrics['post_load_passed']}/"
            f"{metrics['post_load_total']}"
        ),

        (
            "Post-Load Success Rate",
            f"{metrics['post_load_success_rate']:.2f}%"
        ),

        (
            "SQL Validation",
            f"{metrics['sql_passed']}/"
            f"{metrics['sql_total']}"
        ),

        (
            "SQL Success Rate",
            f"{metrics['sql_success_rate']:.2f}%"
        )

    ]


    start_row = 3


    for index, metric in enumerate(
        metrics_data
    ):

        metric_name = metric[0]

        metric_value = metric[1]

        row = (
            start_row + index
        )


        name_cell = worksheet.cell(
            row=row,
            column=1,
            value=metric_name
        )


        value_cell = worksheet.cell(
            row=row,
            column=2,
            value=metric_value
        )


        name_cell.font = Font(
            bold=True
        )


        name_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )


        value_cell.alignment = Alignment(
            horizontal="center"
        )


    worksheet.column_dimensions[
        "A"
    ].width = 30


    worksheet.column_dimensions[
        "B"
    ].width = 20


    # ========================================================
    # SOURCE VS VALID TABLE
    # ========================================================

    worksheet["D3"] = "Table"

    worksheet["E3"] = "Source Records"

    worksheet["F3"] = "Valid Records"


    headers = [

        worksheet["D3"],

        worksheet["E3"],

        worksheet["F3"]

    ]


    for cell in headers:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )


        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )


        cell.alignment = Alignment(
            horizontal="center"
        )


    tables = [

        "customers",

        "accounts",

        "transactions"

    ]


    for index, table in enumerate(
        tables,
        start=4
    ):

        worksheet.cell(
            row=index,
            column=4,
            value=table.title()
        )


        worksheet.cell(
            row=index,
            column=5,
            value=metrics[
                "source_counts"
            ].get(
                table,
                0
            )
        )


        worksheet.cell(
            row=index,
            column=6,
            value=metrics[
                "valid_counts"
            ].get(
                table,
                0
            )
        )


    worksheet.column_dimensions[
        "D"
    ].width = 18


    worksheet.column_dimensions[
        "E"
    ].width = 18


    worksheet.column_dimensions[
        "F"
    ].width = 18


    # ========================================================
    # BAR CHART
    # ========================================================

    bar_chart = BarChart()

    bar_chart.type = "col"

    bar_chart.style = 10

    bar_chart.title = (
        "Source vs Valid Records"
    )

    bar_chart.y_axis.title = (
        "Record Count"
    )

    bar_chart.x_axis.title = (
        "Table"
    )


    chart_data = Reference(
        worksheet,
        min_col=5,
        max_col=6,
        min_row=3,
        max_row=6
    )


    chart_categories = Reference(
        worksheet,
        min_col=4,
        min_row=4,
        max_row=6
    )


    bar_chart.add_data(
        chart_data,
        titles_from_data=True
    )


    bar_chart.set_categories(
        chart_categories
    )


    bar_chart.height = 8

    bar_chart.width = 14


    worksheet.add_chart(
        bar_chart,
        "D9"
    )


    # ========================================================
    # ISSUES BY TABLE
    # ========================================================

    if (
        not data_quality_df.empty
        and "table" in data_quality_df.columns
    ):

        issues_by_table = (
            data_quality_df[
                "table"
            ]
            .value_counts()
        )


        worksheet["H3"] = (
            "Issues by Table"
        )

        worksheet["I3"] = (
            "Issue Count"
        )


        for cell in [

            worksheet["H3"],

            worksheet["I3"]

        ]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )


            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )


        for index, (
            table,
            count
        ) in enumerate(
            issues_by_table.items(),
            start=4
        ):

            worksheet.cell(
                row=index,
                column=8,
                value=table
            )


            worksheet.cell(
                row=index,
                column=9,
                value=count
            )


    worksheet.column_dimensions[
        "H"
    ].width = 20


    worksheet.column_dimensions[
        "I"
    ].width = 15


    # ========================================================
    # PRIORITY SUMMARY
    # ========================================================

    if (
        not rca_df.empty
        and "priority" in rca_df.columns
    ):

        priority_counts = (
            rca_df[
                "priority"
            ]
            .value_counts()
        )


        worksheet["H9"] = "Priority"

        worksheet["I9"] = "Issue Count"


        for cell in [

            worksheet["H9"],

            worksheet["I9"]

        ]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )


            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )


        for index, (
            priority,
            count
        ) in enumerate(
            priority_counts.items(),
            start=10
        ):

            worksheet.cell(
                row=index,
                column=8,
                value=priority
            )


            worksheet.cell(
                row=index,
                column=9,
                value=count
            )


        # ====================================================
        # PIE CHART
        # ====================================================

        pie_chart = PieChart()

        pie_chart.title = (
            "Issue Priority Distribution"
        )


        pie_data = Reference(
            worksheet,
            min_col=9,
            min_row=9,
            max_row=(
                9 +
                len(priority_counts)
            )
        )


        pie_labels = Reference(
            worksheet,
            min_col=8,
            min_row=10,
            max_row=(
                9 +
                len(priority_counts)
            )
        )


        pie_chart.add_data(
            pie_data,
            titles_from_data=True
        )


        pie_chart.set_categories(
            pie_labels
        )


        pie_chart.height = 8

        pie_chart.width = 12


        worksheet.add_chart(
            pie_chart,
            "H15"
        )


# ============================================================
# CREATE EXCEL DASHBOARD
# ============================================================

def create_excel_dashboard():

    print(
        "\nLoading project reports..."
    )


    (
        data_quality_df,
        rca_df,
        post_load_df,
        sql_validation_df
    ) = load_reports()


    print(
        "Calculating dynamic metrics..."
    )


    metrics = calculate_metrics(
        data_quality_df,
        rca_df,
        post_load_df,
        sql_validation_df
    )


    print(
        "Creating Excel workbook..."
    )


    workbook = Workbook()


    # ========================================================
    # DASHBOARD
    # ========================================================

    create_dashboard(
        workbook,
        metrics,
        data_quality_df,
        rca_df
    )


    # ========================================================
    # DATA QUALITY ISSUES
    # ========================================================

    worksheet = workbook.create_sheet(
        "Data Quality Issues"
    )


    add_dataframe_to_sheet(
        worksheet,
        data_quality_df
    )


    # ========================================================
    # ROOT CAUSE ANALYSIS
    # ========================================================

    worksheet = workbook.create_sheet(
        "Root Cause Analysis"
    )


    add_dataframe_to_sheet(
        worksheet,
        rca_df
    )


    # ========================================================
    # POST LOAD VALIDATION
    # ========================================================

    worksheet = workbook.create_sheet(
        "Post Load Validation"
    )


    add_dataframe_to_sheet(
        worksheet,
        post_load_df
    )


    # ========================================================
    # SQL VALIDATION
    # ========================================================

    worksheet = workbook.create_sheet(
        "SQL Validation"
    )


    add_dataframe_to_sheet(
        worksheet,
        sql_validation_df
    )


    # ========================================================
    # SAVE WORKBOOK
    # ========================================================

    workbook.save(
        OUTPUT_FILE
    )


    # ========================================================
    # TERMINAL SUMMARY
    # ========================================================

    print(
        "\n" + "=" * 60
    )


    print(
        "DYNAMIC EXCEL DASHBOARD CREATED SUCCESSFULLY"
    )


    print(
        "=" * 60
    )


    print(
        "\nSOURCE RECORDS"
    )


    print(
        f"Customers: "
        f"{metrics['source_counts'].get('customers', 0)}"
    )


    print(
        f"Accounts: "
        f"{metrics['source_counts'].get('accounts', 0)}"
    )


    print(
        f"Transactions: "
        f"{metrics['source_counts'].get('transactions', 0)}"
    )


    print(
        "\nVALID RECORDS"
    )


    print(
        f"Customers: "
        f"{metrics['valid_counts'].get('customers', 0)}"
    )


    print(
        f"Accounts: "
        f"{metrics['valid_counts'].get('accounts', 0)}"
    )


    print(
        f"Transactions: "
        f"{metrics['valid_counts'].get('transactions', 0)}"
    )


    print(
        "\nMIGRATION METRICS"
    )


    print(
        f"Total Source Records: "
        f"{metrics['total_source_records']}"
    )


    print(
        f"Valid Records: "
        f"{metrics['total_valid_records']}"
    )


    print(
        f"Invalid Records: "
        f"{metrics['total_invalid_records']}"
    )


    print(
        f"Migration Success Rate: "
        f"{metrics['migration_success_rate']:.2f}%"
    )


    print(
        "\nDATA QUALITY METRICS"
    )


    print(
        f"Total Issues: "
        f"{metrics['total_issues']}"
    )


    print(
        f"High Priority Issues: "
        f"{metrics['high_priority']}"
    )


    print(
        f"Medium Priority Issues: "
        f"{metrics['medium_priority']}"
    )


    print(
        f"Low Priority Issues: "
        f"{metrics['low_priority']}"
    )


    print(
        "\nVALIDATION METRICS"
    )


    print(
        f"Post-Load Validation: "
        f"{metrics['post_load_passed']}/"
        f"{metrics['post_load_total']}"
    )


    print(
        f"Post-Load Success Rate: "
        f"{metrics['post_load_success_rate']:.2f}%"
    )


    print(
        f"SQL Validation: "
        f"{metrics['sql_passed']}/"
        f"{metrics['sql_total']}"
    )


    print(
        f"SQL Success Rate: "
        f"{metrics['sql_success_rate']:.2f}%"
    )


    print(
        "\nDashboard location:"
    )


    print(
        OUTPUT_FILE
    )


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    create_excel_dashboard()